"""
Generate comprehensive Word (.docx) screenshot guide for TrustForge.
Fully reconciled with actual codebase implementation (PBKDF2, HMAC-SHA256 tokens,
6 Trust Score pillars, 10 model benchmark slots M01-M10, and 7 primary TOPSIS criteria).
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

CATALOG = [
    {
        "id": "fig_01",
        "file": "fig_01_landing_hero.png",
        "title": "Figure 1: TrustForge Public Landing Page & Privacy Assurance Hero",
        "route": "/ (Public)",
        "role": "Public Visitor",
        "category": "Customer Platform & Identity Workflows",
        "desc": (
            "The landing page provides the primary interface for organizations seeking automated self-assessment against the DPDP Act 2023. "
            "It highlights the platform's core architectural components: hybrid RAG retrieval, multi-model benchmarking, and evidence management. "
            "Visitors can initiate a non-authenticated preliminary domain scan or proceed to create an enterprise account."
        )
    },
    {
        "id": "fig_02",
        "file": "fig_02_public_scan_page.png",
        "title": "Figure 2: Instant Public Domain Assessment & Scanner",
        "route": "/scan (Public)",
        "role": "Public Visitor",
        "category": "Customer Platform & Identity Workflows",
        "desc": (
            "Executes an unauthenticated, non-invasive perimeter scan of a target web domain to check public privacy configurations. "
            "It checks for the presence of cookie banners, privacy policy links, SSL/TLS transport availability, and detectable third-party scripts. "
            "The resulting summary identifies observable perimeter issues and provides recommendations for full compliance auditing."
        )
    },
    {
        "id": "fig_03",
        "file": "fig_03_customer_login_page.png",
        "title": "Figure 3: Customer Authentication Portal",
        "route": "/login (Public/Client)",
        "role": "Client User / Compliance Lead",
        "category": "Customer Platform & Identity Workflows",
        "desc": (
            "Implements multi-tenant credential authentication using standard PBKDF2 password hashing (SHA-512, 10,000 iterations) with unique salt. "
            "Upon successful authentication, the backend issues an HMAC-SHA256 session token with tenant organization claims. "
            "Tenant-scoped route middleware ensures users cannot access or view data belonging to other registered organizations."
        )
    },
    {
        "id": "fig_04",
        "file": "fig_04_customer_signup_page.png",
        "title": "Figure 4: Enterprise Registration & Account Creation",
        "route": "/signup (Public)",
        "role": "Prospective Client",
        "category": "Customer Platform & Identity Workflows",
        "desc": (
            "Allows new enterprise compliance personnel to register a tenant account with email validation and password complexity checks. "
            "Generates a dedicated organization identifier and provisions initial tenant records in the durable application store. "
            "Upon registration, the user is transitioned into the onboarding workflow to configure initial organizational parameters."
        )
    },
    {
        "id": "fig_05",
        "file": "fig_05_onboarding_page.png",
        "title": "Figure 5: Enterprise Organization Profile Setup Wizard",
        "route": "/onboarding (Authenticated Client)",
        "role": "Organization Administrator",
        "category": "Customer Platform & Identity Workflows",
        "desc": (
            "Collects organizational metadata including corporate entity name, industry sector, primary website domain, and designated DPO details. "
            "Determines whether the entity self-classifies as a Significant Data Fiduciary (SDF) under DPDP Act Section 10 obligations. "
            "Initializes default assessment baselines and generates domain verification tokens for ownership validation."
        )
    },
    {
        "id": "fig_06",
        "file": "fig_06_dashboard_overview.png",
        "title": "Figure 6: Executive Privacy Compliance Dashboard",
        "route": "/dashboard (Authenticated Client)",
        "role": "Compliance Officer / DPO",
        "category": "Customer Platform & Identity Workflows",
        "desc": (
            "Displays the tenant's current Trust Score, assessment completion progress, active findings count, and evidence metrics. "
            "Summary cards categorize compliance observations across the six statutory pillars and highlight open remediation items. "
            "Provides direct links to start or resume the 14-step assessment, trigger perimeter scans, and export summary reports."
        )
    },
    {
        "id": "fig_07",
        "file": "fig_07_org_settings_profile.png",
        "title": "Figure 7: Organization Profile & Data Fiduciary Configuration",
        "route": "/settings/organization (Client Admin)",
        "role": "Organization Administrator",
        "category": "Customer Platform & Identity Workflows",
        "desc": (
            "Allows administrators to update organizational records, designated Data Protection Officer contact details, and corporate headquarters. "
            "Enables maintenance of corporate sector classifications and review of applicable DPDP Act statutory requirements. "
            "All profile updates are recorded in the organization audit history with timestamps and user identifiers."
        )
    },
    {
        "id": "fig_08",
        "file": "fig_08_org_settings_dns_challenge.png",
        "title": "Figure 8: Automated DNS TXT Domain Ownership Verification",
        "route": "/settings/organization#domain-verification",
        "role": "Organization Administrator",
        "category": "Customer Platform & Identity Workflows",
        "desc": (
            "Generates a unique DNS TXT challenge token formatted as `trustforge-site-verification=<token>` to verify domain ownership. "
            "The backend performs asynchronous DNS lookups against public nameservers to confirm token propagation prior to active scanning. "
            "Ensures website scans and compliance scoring are only initiated for domains legitimately managed by the tenant organization."
        )
    },
    {
        "id": "fig_09",
        "file": "fig_09_assessment_step01_governance.png",
        "title": "Figure 9: Assessment Step 1 — Governance & Accountability Architecture",
        "route": "/assessment/wizard/step-1",
        "role": "Compliance Assessor",
        "category": "14-Step Statutory Assessment Wizard",
        "desc": (
            "Evaluates organizational data governance mechanisms, internal compliance reporting hierarchies, and DPO appointments. "
            "Maps directly to DPDP Act Section 8(1) and Section 10(2)(a) obligations for responsible data fiduciary governance. "
            "Audits formal documentation, board-level oversight policies, and designated privacy officer operational roles."
        )
    },
    {
        "id": "fig_10",
        "file": "fig_10_assessment_step02_consent.png",
        "title": "Figure 10: Assessment Step 2 — Consent Architecture & Validation",
        "route": "/assessment/wizard/step-2",
        "role": "Compliance Assessor",
        "category": "14-Step Statutory Assessment Wizard",
        "desc": (
            "Audits affirmative, freely given, specific, informed, and unambiguous consent mechanisms under DPDP Act Section 6(1). "
            "Checks that opt-in selections are unchecked by default, bundled terms are avoided, and withdrawal procedures are accessible. "
            "Verifies timestamped consent recordkeeping and specific purpose declarations for each data processing activity."
        )
    },
    {
        "id": "fig_11",
        "file": "fig_11_assessment_step03_notice.png",
        "title": "Figure 11: Assessment Step 3 — Multilingual Statutory Notice Delivery",
        "route": "/assessment/wizard/step-3",
        "role": "Compliance Assessor",
        "category": "14-Step Statutory Assessment Wizard",
        "desc": (
            "Evaluates statutory privacy notice completeness, itemization of collected personal data, and DPO contact details under Section 5. "
            "Checks support for notice availability in English and the 22 languages specified in the Eighth Schedule of the Constitution of India. "
            "Verifies clear instructions on how data principals may exercise rights or lodge grievances with the fiduciary."
        )
    },
    {
        "id": "fig_12",
        "file": "fig_12_assessment_step04_consent_mgr.png",
        "title": "Figure 12: Assessment Step 4 — Consent Manager Integration & Interoperability",
        "route": "/assessment/wizard/step-4",
        "role": "Compliance Assessor",
        "category": "14-Step Statutory Assessment Wizard",
        "desc": (
            "Assesses technical readiness for integrating with Data Protection Board-registered Consent Managers under Section 6(7)-(9). "
            "Evaluates API interfaces and workflows enabling data principals to view, manage, and withdraw consent through unified mechanisms. "
            "Audits procedures for propagating consent revocation across internal databases and third-party data processors."
        )
    },
    {
        "id": "fig_13",
        "file": "fig_13_assessment_step05_minimization.png",
        "title": "Figure 13: Assessment Step 5 — Purpose Limitation & Data Minimization",
        "route": "/assessment/wizard/step-5",
        "role": "Compliance Assessor",
        "category": "14-Step Statutory Assessment Wizard",
        "desc": (
            "Assesses adherence to statutory collection boundaries under DPDP Act Section 4 and Section 7, ensuring only necessary data is collected. "
            "Identifies potential collection creep, collateral data gathering, and secondary use beyond the consented purpose. "
            "Checks safeguards preventing personal data collected for transactional purposes from being reused for unauthorized profiling."
        )
    },
    {
        "id": "fig_14",
        "file": "fig_14_assessment_step06_retention.png",
        "title": "Figure 14: Assessment Step 6 — Storage Limitation & Secure Erasure",
        "route": "/assessment/wizard/step-6",
        "role": "Compliance Assessor",
        "category": "14-Step Statutory Assessment Wizard",
        "desc": (
            "Evaluates defined data retention schedules, automated archival triggers, and secure deletion protocols under Section 8(7). "
            "Verifies that personal data is deleted once the specified processing purpose is accomplished or consent is formally revoked. "
            "Audits tombstone tracking records, deletion verification logs, and legal hold override controls."
        )
    },
    {
        "id": "fig_15",
        "file": "fig_15_assessment_step07_accuracy.png",
        "title": "Figure 15: Assessment Step 7 — Data Quality, Completeness & Accuracy",
        "route": "/assessment/wizard/step-7",
        "role": "Compliance Assessor",
        "category": "14-Step Statutory Assessment Wizard",
        "desc": (
            "Assesses mechanisms for maintaining data accuracy and completeness when personal data is used to make decisions affecting data principals. "
            "Maps directly to DPDP Act Section 8(3) obligations requiring reasonable measures to ensure data consistency. "
            "Reviews user self-service correction interfaces, synchronization processes, and downstream processor notification workflows."
        )
    },
    {
        "id": "fig_16",
        "file": "fig_16_assessment_step08_safeguards.png",
        "title": "Figure 16: Assessment Step 8 — Reasonable Technical & Security Safeguards",
        "route": "/assessment/wizard/step-8",
        "role": "Compliance Assessor",
        "category": "14-Step Statutory Assessment Wizard",
        "desc": (
            "Audits technical security controls including TLS encryption in transit, role-based access control, and credential safeguards. "
            "Evaluates organizational adherence to DPDP Act Section 8(5) requiring reasonable security safeguards to prevent data breaches. "
            "Reviews vulnerability review cadences, system patching policies, and multi-factor authentication enforcement."
        )
    },
    {
        "id": "fig_17",
        "file": "fig_17_assessment_step09_breach.png",
        "title": "Figure 17: Assessment Step 9 — Personal Data Breach Incident Management",
        "route": "/assessment/wizard/step-9",
        "role": "Compliance Assessor",
        "category": "14-Step Statutory Assessment Wizard",
        "desc": (
            "Reviews incident response procedures and formal breach reporting plans as required under DPDP Act Section 8(6). "
            "Evaluates playbooks for notifying the Data Protection Board of India and affected data principals upon discovering a personal data breach. "
            "Checks incident containment steps, forensic audit logging mechanisms, and statutory reporting escalation timelines."
        )
    },
    {
        "id": "fig_18",
        "file": "fig_18_assessment_step10_children.png",
        "title": "Figure 18: Assessment Step 10 — Children's Data & Persons with Disabilities",
        "route": "/assessment/wizard/step-10",
        "role": "Compliance Assessor",
        "category": "14-Step Statutory Assessment Wizard",
        "desc": (
            "Evaluates parental consent verification protocols and strict prohibitions against tracking, profiling, or targeted advertising to minors. "
            "Enforces compliance with DPDP Act Section 9 mandates protecting children and persons under legal guardianship. "
            "Reviews age-verification workflows, neutral gating methods, and absence of harmful behavioral monitoring practices."
        )
    },
    {
        "id": "fig_19",
        "file": "fig_19_assessment_step11_rights.png",
        "title": "Figure 19: Assessment Step 11 — Data Principal Rights Redressal Architecture",
        "route": "/assessment/wizard/step-11",
        "role": "Compliance Assessor",
        "category": "14-Step Statutory Assessment Wizard",
        "desc": (
            "Audits self-service procedures for processing data principal requests: right to access (Sec 11), correction (Sec 12), and erasure. "
            "Evaluates identity verification methods, data export formats, and response tracking against statutory fulfillment windows. "
            "Reviews mechanisms supporting the right of nomination in the event of death or legal incapacity under Section 14."
        )
    },
    {
        "id": "fig_20",
        "file": "fig_20_assessment_step12_grievance.png",
        "title": "Figure 20: Assessment Step 12 — Grievance Redressal Mechanism & Timelines",
        "route": "/assessment/wizard/step-12",
        "role": "Compliance Assessor",
        "category": "14-Step Statutory Assessment Wizard",
        "desc": (
            "Evaluates internal grievance redressal procedures and Data Protection Officer escalation channels under Section 13. "
            "Verifies ticket logging systems that record complaints and track resolution turnaround times. "
            "Checks that communications inform data principals of their right to appeal unresolved grievances to the Data Protection Board."
        )
    },
    {
        "id": "fig_21",
        "file": "fig_21_assessment_step13_sdf.png",
        "title": "Figure 21: Assessment Step 13 — Significant Data Fiduciary (SDF) Obligations",
        "route": "/assessment/wizard/step-13",
        "role": "Compliance Assessor",
        "category": "14-Step Statutory Assessment Wizard",
        "desc": (
            "Audits higher-tier obligations applicable to designated Significant Data Fiduciaries under DPDP Act Section 10. "
            "Assesses periodic Data Protection Impact Assessment (DPIA) frameworks, independent audit readiness, and resident DPO appointments. "
            "Reviews algorithmic risk monitoring procedures and board reporting cadences for systemic privacy impacts."
        )
    },
    {
        "id": "fig_22",
        "file": "fig_22_assessment_step14_review.png",
        "title": "Figure 22: Assessment Step 14 — Final Audit Review & Evidence Synthesis",
        "route": "/assessment/wizard/step-14",
        "role": "Compliance Assessor",
        "category": "14-Step Statutory Assessment Wizard",
        "desc": (
            "Consolidates responses from all 14 statutory control sections into a pre-submission review matrix with completeness indicators. "
            "Allows the compliance assessor to inspect consistency, attach supplementary documentation, and confirm submission accuracy. "
            "Final submission locks the assessment version and initiates deterministic scoring and RAG citation validation."
        )
    },
    {
        "id": "fig_23",
        "file": "fig_23_assessment_results.png",
        "title": "Figure 23: Complete Statutory Compliance Evaluation Results",
        "route": "/assessment/results/:id",
        "role": "Compliance Officer / DPO",
        "category": "14-Step Statutory Assessment Wizard",
        "desc": (
            "Presents the completed audit scorecard, section-by-section compliance percentages, and identified statutory gap summaries. "
            "Categorizes compliance gaps by severity (High, Medium, Low) based on statutory obligations and potential exposure tiers. "
            "Generates actionable remediation recommendations and provides access to printable summary reports."
        )
    },
    {
        "id": "fig_24",
        "file": "fig_24_trust_score_page.png",
        "title": "Figure 24: TrustForge Composite Trust Score & Empirical Analytics",
        "route": "/trust-score",
        "role": "Executive / Auditor",
        "category": "Compliance Engines & Analytics",
        "desc": (
            "Displays the deterministic Trust Score computed across six statutory pillars: Privacy, Security, Transparency, Governance, Lifecycle, Rights. "
            "Incorporates faculty C1–C10 compliance mappings, dual weighting schemes (W1 Penalty-Proportional, W2 Equal), and posture bands. "
            "Enforces critical gating caps (capping scores at 40 if HTTPS or privacy notices are missing) and provides pillar radar visualizations."
        )
    },
    {
        "id": "fig_25",
        "file": "fig_25_website_scan_page.png",
        "title": "Figure 25: Automated Perimeter Security & Privacy Cookie Scanner",
        "route": "/scans",
        "role": "Security Engineer / Auditor",
        "category": "Compliance Engines & Analytics",
        "desc": (
            "Executes automated analysis of verified enterprise web properties to identify tracking cookies and missing security headers. "
            "Evaluates SSL/TLS transport availability, defensive HTTP headers (HSTS, CSP), and accessibility of privacy notices. "
            "Observed perimeter findings are automatically linked into the findings vault for remediation tracking."
        )
    },
    {
        "id": "fig_26",
        "file": "fig_26_findings_page.png",
        "title": "Figure 26: Statutory Vulnerabilities & Findings Management Vault",
        "route": "/findings",
        "role": "Compliance Manager",
        "category": "Compliance Engines & Analytics",
        "desc": (
            "Centralizes compliance gaps and perimeter scan findings in a searchable issue tracker categorized by statutory section and severity. "
            "Links each item to relevant DPDP Act provisions, descriptive remediation notes, and current resolution status. "
            "Allows compliance managers to assign remediation owners and update remediation progress."
        )
    },
    {
        "id": "fig_27",
        "file": "fig_27_evidence_page.png",
        "title": "Figure 27: Cryptographic Evidence Locker & Document Management",
        "route": "/evidence",
        "role": "Compliance Officer",
        "category": "Compliance Engines & Analytics",
        "desc": (
            "Provides an auditable evidence locker where organizations upload policy documents, consent logs, and architecture documentation. "
            "Each uploaded file is hashed using SHA-256 upon ingestion to establish a tamper-evident chain of custody and file verification record. "
            "Stored evidence is indexed for retrieval to support citation grounding during automated compliance assessments."
        )
    },
    {
        "id": "fig_28",
        "file": "fig_28_remediation_page.png",
        "title": "Figure 28: Interactive Statutory Remediation Roadmaps & Workflows",
        "route": "/remediation",
        "role": "Remediation Lead / Engineer",
        "category": "Compliance Engines & Analytics",
        "desc": (
            "Translates compliance findings into an interactive task-based engineering and policy remediation roadmap. "
            "Tracks implementation progress across technical controls, policy updates, and operational procedure revisions. "
            "Provides contextual guidance and statutory references to assist internal teams in closing identified gaps."
        )
    },
    {
        "id": "fig_29",
        "file": "fig_29_reports_page.png",
        "title": "Figure 29: Statutory Audit Report Generation & Export Console",
        "route": "/reports",
        "role": "Auditor / Executive",
        "category": "Compliance Engines & Analytics",
        "desc": (
            "Enables export of internal compliance assessment reports in PDF and structured JSON data formats for governance review. "
            "Includes executive summaries, pillar score breakdowns, statutory section citations, and open finding inventories. "
            "Allows filtering reports by assessment date and scope for internal stakeholder review and audit preparation."
        )
    },
    {
        "id": "fig_30",
        "file": "fig_30_model_benchmark_page.png",
        "title": "Figure 30: Multi-Model Evaluation & TOPSIS Decision Matrix",
        "route": "/benchmark",
        "role": "AI Research / DPO",
        "category": "Compliance Engines & Analytics",
        "desc": (
            "Presents empirical benchmark metrics across 10 designated candidate model slots (M01–M10) evaluated against DPDP Act controls. "
            "Applies the TOPSIS multi-criteria decision algorithm using 7 primary criteria (Macro-F1, Citations, Evidence, Consistency, Latency, Cost, Feasibility). "
            "Maintains zero-paid commercial API spend ($0.00) by utilizing local models and free tiers, gating billable providers M06–M09."
        )
    },
    {
        "id": "fig_31",
        "file": "fig_31_rag_monitor_page.png",
        "title": "Figure 31: Real-Time RAG Diagnostic & Retrieval Fusion Monitor",
        "route": "/rag-monitor",
        "role": "Data Engineer / Auditor",
        "category": "Compliance Engines & Analytics",
        "desc": (
            "Provides real-time visibility into the hybrid dense/sparse retrieval pipeline combining vector search with lexical BM25. "
            "Monitors Reciprocal Rank Fusion (RRF, k=60), cosine similarity thresholds, and statutory corpus chunk boundaries. "
            "Enables engineers to verify that model findings are grounded in retrieved excerpts from the DPDP Act 2023 legal corpus."
        )
    },
    {
        "id": "fig_32",
        "file": "fig_32_rights_portal_page.png",
        "title": "Figure 32: Data Principal Rights Request Management Portal",
        "route": "/rights-portal",
        "role": "Privacy Operations Lead",
        "category": "Compliance Engines & Analytics",
        "desc": (
            "Manages incoming data principal privacy requests including Access, Correction, Erasure, and Grievance submissions. "
            "Tracks operational turnaround times against statutory 30-day response targets with visual status badges. "
            "Records request processing milestones and documentation to maintain a complete internal fulfillment history."
        )
    },
    {
        "id": "fig_33",
        "file": "fig_33_admin_login_page.png",
        "title": "Figure 33: Dedicated Administrative Authentication Gate",
        "route": "/admin/login (Dedicated)",
        "role": "System Administrator",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Enforces an isolated administrative authentication gate requiring explicit server-verified platform administrator roles. "
            "Client customer credentials submitted here are rejected by the backend to prevent unauthorized console access. "
            "Requires a non-empty password and verifies the user profile dynamically against database authority before granting access."
        )
    },
    {
        "id": "fig_34",
        "file": "fig_34_admin_login_error.png",
        "title": "Figure 34: Administrative Access Control & Privilege Rejection Error",
        "route": "/admin/login (Failure State)",
        "role": "Unauthorized User",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Demonstrates the active access control rejection boundary when invalid or non-administrative credentials attempt console login. "
            "Displays a clear, non-leaking authentication error message while recording the unauthorized attempt in the security audit log. "
            "Prevents privilege escalation and enforces strict separation between client tenant roles and platform administrators."
        )
    },
    {
        "id": "fig_35",
        "file": "fig_35_admin_command_center.png",
        "title": "Figure 35: Central Platform Command Center & Global KPIs",
        "route": "/admin/dashboard",
        "role": "Platform Superadmin",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Aggregates platform-wide telemetry, cross-tenant compliance indices, active scan queues, and aggregate assessment totals. "
            "Provides platform administrators with operational situational awareness regarding system activity and workload distribution. "
            "Displays overview charts of system throughput, scan processing times, and aggregate finding categories."
        )
    },
    {
        "id": "fig_36",
        "file": "fig_36_admin_analytics_page.png",
        "title": "Figure 36: Cross-Tenant Compliance Analytics & Distribution Trends",
        "route": "/admin/analytics",
        "role": "Compliance Analyst",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Presents aggregated cross-organization compliance trends, common statutory gap areas, and cross-sector comparisons. "
            "Visualizes the distribution of Trust Scores across industry categories (Healthcare, Finance, SaaS, Retail). "
            "Enables platform analysts to identify recurring compliance bottlenecks under the DPDP Act across registered tenants."
        )
    },
    {
        "id": "fig_37",
        "file": "fig_37_admin_organizations_page.png",
        "title": "Figure 37: Multi-Tenant Enterprise Organization Management",
        "route": "/admin/organizations",
        "role": "Platform Superadmin",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Manages registered tenant organizations, subscription tier settings, domain verification statuses, and compliance postures. "
            "Provides administrative controls to update tenant status, configure resource allocations, and inspect organization profiles. "
            "Supports searching and filtering across registered legal entities, verification flags, and Significant Data Fiduciary designations."
        )
    },
    {
        "id": "fig_38",
        "file": "fig_38_admin_organization_detail.png",
        "title": "Figure 38: Granular Organization Tenant Governance & Deep Inspection",
        "route": "/admin/organizations/:id",
        "role": "Platform Superadmin",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Provides detailed inspection of a specific tenant's settings, assigned DPO details, DNS challenge records, and audit history. "
            "Enables manual domain verification controls for authorized enterprise testing or air-gapped organizational environments. "
            "Displays tenant-specific assessment history, registered members, and recent compliance milestones."
        )
    },
    {
        "id": "fig_39",
        "file": "fig_39_admin_users_page.png",
        "title": "Figure 39: Enterprise User Directory & Role Assignment Controls",
        "route": "/admin/users",
        "role": "Security Administrator",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Centralizes platform user management across registered tenants, displaying assigned roles and account creation dates. "
            "Enables administrators to create user accounts, trigger password resets, and adjust role assignments (Admin, DPO, Auditor). "
            "Enforces role-based separation of responsibilities and records all administrative account changes in the audit ledger."
        )
    },
    {
        "id": "fig_40",
        "file": "fig_40_admin_user_detail.png",
        "title": "Figure 40: User Profile Deep Inspection & Session Revocation",
        "route": "/admin/users/:id",
        "role": "Security Administrator",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Provides detailed view of individual user accounts, including associated tenant organizations, role scopes, and recent activity. "
            "Allows administrators to revoke active sessions in the event of credential rotation or employee offboarding. "
            "Audits user authentication events, account status modifications, and session lifetimes."
        )
    },
    {
        "id": "fig_41",
        "file": "fig_41_admin_assessments_page.png",
        "title": "Figure 41: Platform-Wide Assessment Governance & Audit Logs",
        "route": "/admin/assessments",
        "role": "Supervisory Auditor",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Displays all statutory compliance assessments completed across the platform, including completion dates, scores, and review status. "
            "Enables supervisory auditors to review submitted assessment questionnaires, inspect attached evidence, and leave review notes. "
            "Allows cross-tenant compliance comparisons and drill-down into specific assessment responses."
        )
    },
    {
        "id": "fig_42",
        "file": "fig_42_admin_scans_page.png",
        "title": "Figure 42: Global Perimeter Scan Dispatcher & Queue Telemetry",
        "route": "/admin/scans",
        "role": "Systems Administrator",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Monitors the background scan processing queue, showing active domain crawl tasks, completion times, and error rates. "
            "Provides administrative controls to re-run failed scan jobs and inspect scanner worker status. "
            "Logs HTTP status codes, SSL certificate parameters, and detected tracker counts for each completed scan."
        )
    },
    {
        "id": "fig_43",
        "file": "fig_43_admin_findings_page.png",
        "title": "Figure 43: Platform-Wide Vulnerability & Defect Registry",
        "route": "/admin/findings",
        "role": "Chief Compliance Officer",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Consolidates statutory non-compliance observations and perimeter findings across all registered organizations into one view. "
            "Allows filtering by severity level (Critical, High, Medium, Low), DPDP Act section, and current resolution status. "
            "Assists compliance leads in identifying recurring systemic issues and monitoring remediation turnaround across tenants."
        )
    },
    {
        "id": "fig_44",
        "file": "fig_44_admin_evidence_page.png",
        "title": "Figure 44: Global Cryptographic Evidence Audit & Verification Vault",
        "route": "/admin/evidence",
        "role": "Forensic Auditor",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Enables platform auditors to review evidence files uploaded by tenants along with their SHA-256 cryptographic hashes. "
            "Displays file metadata, uploading tenant identity, ingestion timestamps, and linked statutory assessment sections. "
            "Verifies that uploaded documentary evidence maintains data integrity without modification post-submission."
        )
    },
    {
        "id": "fig_45",
        "file": "fig_45_admin_review_page.png",
        "title": "Figure 45: Formal Compliance Review & Certification Authority",
        "route": "/admin/reviews",
        "role": "Senior Certification Officer",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Provides an administrative interface where compliance reviewers evaluate and certify completed tenant assessment submissions. "
            "Allows the reviewer to approve, request revisions on, or reject submitted assessments with itemized feedback. "
            "Approved reviews record an administrative certification timestamp and update the organization's official audit record."
        )
    },
    {
        "id": "fig_46",
        "file": "fig_46_admin_rules_page.png",
        "title": "Figure 46: Statutory Rules Engine & Legal Weight Tuning Matrix",
        "route": "/admin/rules",
        "role": "Legal / Rule Engineer",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Allows authorized administrators to adjust statutory scoring formulas, section weights, and penalty parameters. "
            "Enables updates to assessment rule configurations as official notifications and subordinate rules are issued under the DPDP Act. "
            "Maintains versioned rule profiles to ensure audit assessments remain aligned with prevailing regulatory guidance."
        )
    },
    {
        "id": "fig_47",
        "file": "fig_47_admin_reports_page.png",
        "title": "Figure 47: Global Regulatory Export & Supervisory Report Vault",
        "route": "/admin/reports",
        "role": "Supervisory Auditor",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Maintains an archive of generated assessment summary reports, audit extracts, and regulatory disclosure packages. "
            "Enables administrators to view previously generated reports and review document generation timestamps. "
            "Maintains an audit record of report generation and download events for accountability purposes."
        )
    },
    {
        "id": "fig_48",
        "file": "fig_48_admin_system_health.png",
        "title": "Figure 48: System Health Diagnostics & Infrastructure Monitoring",
        "route": "/admin/health",
        "role": "DevOps / Infrastructure Lead",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Monitors application service availability including Express API responsiveness, Python RAG worker status, and database state. "
            "Tracks memory usage, process uptime, active database connections, and background worker queue health. "
            "Provides diagnostic indicators to ensure system components operate within normal performance tolerances."
        )
    },
    {
        "id": "fig_49",
        "file": "fig_49_admin_audit_logs.png",
        "title": "Figure 49: Immutable Security Audit Log & Forensic Activity Trail",
        "route": "/admin/audit",
        "role": "Forensic / Security Auditor",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Maintains an append-only audit trail recording user logins, role changes, assessment submissions, and administrative updates. "
            "Each log entry records UTC timestamp, actor IP address, action performed, target resource, and result status. "
            "Provides an unalterable activity history for internal accountability and compliance verification."
        )
    },
    {
        "id": "fig_50",
        "file": "fig_50_admin_settings_page.png",
        "title": "Figure 50: Platform Global Settings & Policy Engine Configuration",
        "route": "/admin/settings",
        "role": "Platform Superadmin",
        "category": "Administrative Governance & Platform Audit",
        "desc": (
            "Configures platform security parameters including session timeout durations, CORS origin restrictions, and password policies. "
            "Manages automated maintenance schedules, logging verbosity, and zero-paid inference enforcement policies. "
            "Ensures consistent operational policies and security boundaries across the entire application deployment."
        )
    }
]

def generate_docx():
    screenshots_dir = os.path.abspath(r"TrustForge\TrustDocumentation\complete_documentation_screenshots")
    docx_path = os.path.abspath(r"TrustForge\TrustDocumentation\TrustForge_Complete_System_Screenshots_Guide.docx")

    doc = Document()

    # Set 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Document Header / Title
    title = doc.add_heading("TrustForge™ — Complete System Screenshots & Feature Verification Guide", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.runs[0]
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)  # Navy Blue

    subtitle = doc.add_paragraph("Exhaustive Figure-by-Figure Visual & Architectural Verification Catalog")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.runs[0]
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    # Metadata Table
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Platform Version:", "TrustForge v1.0.0-release (Production Golden Master)"),
        ("Verified Release Commit:", "035429f6c76d7ab6d14f2053b3e9d2050f493672"),
        ("Statutory Foundation:", "Digital Personal Data Protection (DPDP) Act 2023 [Act No. 22 of 2023, India]"),
        ("Authentication & Security:", "PBKDF2 Password Hashing (SHA-512, 10k iter), HMAC-SHA256 Session Tokens, SHA-256 Evidence Hashes"),
        ("Trust Score Architecture:", "6 Canonical Pillars, C1–C10 Compliance Taxonomy, W1/W2 Weights, Gating Caps (min(TS, 40))"),
        ("Benchmarking & TOPSIS:", "10 Model Slots (M01–M10), 7 Primary TOPSIS Criteria, Zero-Paid Inference Policy ($0.00 spend)"),
    ]
    for i, (k, v) in enumerate(meta_data):
        row = table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.text = k
        c1.text = v
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.size = Pt(9.5)
        c1.paragraphs[0].runs[0].font.size = Pt(9.5)

    doc.add_paragraph()  # spacing

    intro = doc.add_paragraph(
        "This document provides the definitive, figure-by-figure photographic and architectural record of the TrustForge "
        "compliance auditing and benchmarking platform. Every screen, form, wizard step, analytics dashboard, and administrative "
        "governance console has been rendered from active application execution and captured in deterministic headless automation. "
        "Each figure is paired with a verified explanation describing its user-facing role, statutory DPDP Act grounding, "
        "and technical implementation status in accordance with the released source code."
    )
    intro.runs[0].font.size = Pt(10.5)

    current_cat = None

    for item in CATALOG:
        # Category header if changed
        if item["category"] != current_cat:
            current_cat = item["category"]
            cat_head = doc.add_heading(f"Section: {current_cat}", level=1)
            cat_run = cat_head.runs[0]
            cat_run.font.name = "Calibri"
            cat_run.font.size = Pt(15)
            cat_run.font.bold = True
            cat_run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        # Figure Title
        fig_head = doc.add_heading(item["title"], level=2)
        fig_run = fig_head.runs[0]
        fig_run.font.name = "Calibri"
        fig_run.font.size = Pt(12)
        fig_run.font.bold = True
        fig_run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

        # Badges
        badge_p = doc.add_paragraph()
        badge_run1 = badge_p.add_run(f"Route: {item['route']}  |  Role: {item['role']}")
        badge_run1.font.bold = True
        badge_run1.font.size = Pt(9.5)
        badge_run1.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        # Image
        img_path = os.path.join(screenshots_dir, item["file"])
        if os.path.exists(img_path):
            doc.add_picture(img_path, width=Inches(6.5))
            last_p = doc.paragraphs[-1]
            last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p_err = doc.add_paragraph(f"[Image file not found: {item['file']}]")
            p_err.runs[0].font.color.rgb = RGBColor(0xDC, 0x26, 0x26)

        # Description
        desc_p = doc.add_paragraph(item["desc"])
        desc_p.runs[0].font.size = Pt(10)
        desc_p.runs[0].font.name = "Calibri"

        # Spacer
        doc.add_paragraph()

    doc.save(docx_path)
    print(f"DOCX created successfully at: {docx_path}")

if __name__ == "__main__":
    generate_docx()
